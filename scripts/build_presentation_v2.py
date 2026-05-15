"""
QueryMind AI — Capstone Presentation v2
8 slides with speaker notes.
Run: pyenv shell 3.12.13 && python3 scripts/build_presentation_v2.py
Output: QueryMind_Presentation_v2.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree
import copy
import os

# ── Brand colours ─────────────────────────────────────────────────────────────
SKY    = RGBColor(0x0E, 0xA5, 0xE9)
DARK   = RGBColor(0x0F, 0x17, 0x2A)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GRAY   = RGBColor(0x94, 0xA3, 0xB8)
GREEN  = RGBColor(0x22, 0xC5, 0x5E)
YELLOW = RGBColor(0xF5, 0x9E, 0x0B)
RED    = RGBColor(0xEF, 0x44, 0x44)
PURPLE = RGBColor(0x8B, 0x5C, 0xF6)
ORANGE = RGBColor(0xF9, 0x73, 0x16)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ── Helpers ───────────────────────────────────────────────────────────────────
def add_slide():
    return prs.slides.add_slide(BLANK)

def bg(slide, color=DARK):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def box(slide, left, top, width, height, fill_color=None,
        line_color=None, line_width=Pt(0)):
    shape = slide.shapes.add_shape(
        1, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.line.width = line_width
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
    return shape

def txt(slide, text, left, top, width, height,
        size=16, bold=False, color=WHITE, align=PP_ALIGN.LEFT,
        italic=False):
    txb = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height))
    txb.word_wrap = True
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size   = Pt(size)
    run.font.bold   = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return txb

def bullet_block(slide, items, left, top, width, height,
                 size=15, color=WHITE, bullet="▸  ", line_space=Pt(6)):
    txb = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height))
    txb.word_wrap = True
    tf = txb.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = line_space
        run = p.add_run()
        run.text = f"{bullet}{item}"
        run.font.size = Pt(size)
        run.font.color.rgb = color

def add_notes(slide, notes_text):
    """Add speaker notes to a slide."""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text

def accent_bar(slide, top=0.5, height=0.06):
    box(slide, 0, top, 13.33, height, fill_color=SKY)

def slide_heading(slide, title, subtitle=None):
    accent_bar(slide)
    txt(slide, title, 0.5, 0.65, 12, 0.75,
        size=30, bold=True, color=WHITE)
    if subtitle:
        txt(slide, subtitle, 0.5, 1.35, 12, 0.4,
            size=14, color=GRAY, italic=True)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — Title
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide(); bg(s)
box(s, 0, 0, 13.33, 0.5,  fill_color=SKY)
box(s, 0, 7.0, 13.33, 0.5, fill_color=SKY)

txt(s, "QueryMind AI", 1, 1.5, 11, 1.3,
    size=54, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(s, "Text-to-SQL Analytics Engine", 1, 2.85, 11, 0.65,
    size=26, bold=True, color=SKY, align=PP_ALIGN.CENTER)
txt(s, "Codebasics AI Engineering Bootcamp  •  Assignment 3  •  Capstone Project",
    1, 3.6, 11, 0.45, size=14, color=GRAY, align=PP_ALIGN.CENTER)

box(s, 3.5, 4.3, 6.33, 0.02, fill_color=RGBColor(0x1E, 0x3A, 0x5F))

txt(s, "Vijay Mamilla", 1, 4.5, 11, 0.4,
    size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(s, "Associate Director | AI/ML Engineering Lead | SS&C Technologies",
    1, 4.9, 11, 0.4, size=13, color=GRAY, align=PP_ALIGN.CENTER)
txt(s, "github.com/vijaymamilla/querymind  •  youtu.be/L4pgpwKwABo",
    1, 5.4, 11, 0.4, size=13, color=SKY, align=PP_ALIGN.CENTER)

add_notes(s, """SLIDE 1 — TITLE (30 seconds)

"Hi, I'm Vijay Mamilla — Associate Director and AI/ML Engineering Lead at SS&C Technologies.

This is QueryMind AI — my capstone project for the Codebasics AI Engineering Bootcamp.

QueryMind is a Text-to-SQL Analytics Engine. It converts plain English questions into SQL queries and runs them against a real PostgreSQL database — no SQL knowledge needed.

Let me walk you through what I built, why I built it, and show you a live demo."
""")


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — Problem Statement
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide(); bg(s)
slide_heading(s, "Problem Statement", "Why does QueryMind exist?")

# Left — Problem
box(s, 0.4, 1.85, 5.9, 5.1, fill_color=RGBColor(0x1E, 0x29, 0x3B))
txt(s, "The Problem", 0.7, 1.95, 5.3, 0.45,
    size=16, bold=True, color=RGBColor(0xF8, 0x71, 0x71))
bullet_block(s, [
    "Business analysts need data but can't write SQL",
    "Data teams bottlenecked by ad-hoc query requests",
    "BI tools need pre-built dashboards — no flexibility",
    "Existing NL interfaces lack guardrails & auditability",
    "No visibility into what SQL was generated or why",
], 0.7, 2.45, 5.3, 4.2, size=14,
   color=RGBColor(0xF8, 0xD7, 0xDA), bullet="✗  ")

# Right — Solution
box(s, 6.9, 1.85, 5.9, 5.1, fill_color=RGBColor(0x0C, 0x2A, 0x1F))
txt(s, "QueryMind Solves It", 7.2, 1.95, 5.3, 0.45,
    size=16, bold=True, color=GREEN)
bullet_block(s, [
    "Type plain English → get SQL + results in ~3 seconds",
    "Runtime guardrails block dangerous operations (DELETE)",
    "Every query logged with SQL, status, and latency",
    "Admin panel to tune behavior without touching code",
    "Self-healing SQL retries automatically on error",
], 7.2, 2.45, 5.3, 4.2, size=14,
   color=RGBColor(0xD1, 0xFA, 0xE5), bullet="✓  ")

add_notes(s, """SLIDE 2 — PROBLEM STATEMENT (45 seconds)

"The problem I'm solving is real and common.

Business analysts need data insights — but they can't write SQL. That creates a bottleneck where data teams spend half their time answering ad-hoc query requests instead of doing real analysis.

Traditional BI tools like Tableau require pre-built dashboards — they're not flexible for open-ended questions. And existing natural language interfaces lack guardrails — you can't just let users run arbitrary SQL against production data without any controls.

QueryMind solves all of this:
- Type a plain English question and get results in about 3 seconds
- Runtime guardrails prevent dangerous operations like DELETE
- Every query is fully logged with the generated SQL and latency
- Everything is configurable through an admin UI — no code changes needed
- If the SQL fails, it automatically retries with the error fed back to the LLM"
""")


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — Architecture Diagram
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide(); bg(s)
slide_heading(s, "Architecture — 7-Stage Modular Pipeline",
              "Natural language → validated SQL → formatted results")

# Embed the architecture PNG if it exists
img_path = "QueryMind_Architecture.png"
if os.path.exists(img_path):
    s.shapes.add_picture(img_path,
                         Inches(0.3), Inches(1.6),
                         Inches(12.73), Inches(5.7))
else:
    txt(s, "[ QueryMind_Architecture.png ]", 0.3, 3.5, 12.73, 1.0,
        size=16, color=GRAY, align=PP_ALIGN.CENTER)
    txt(s, "Run scripts/build_diagram.py to generate the image first.",
        0.3, 4.5, 12.73, 0.5, size=12, color=GRAY, align=PP_ALIGN.CENTER)

add_notes(s, """SLIDE 3 — ARCHITECTURE (2 minutes)

"The core of QueryMind is a 7-stage modular pipeline. Let me walk through each stage.

Stage 1 — Classifier: Detects the intent of the question. Is it a simple SELECT, an aggregate with GROUP BY, a multi-table JOIN, a temporal query, or something that should be blocked entirely?

Stage 2 — Schema Linker: Maps the entities mentioned in the question to real table and column names in the PostgreSQL database. So 'revenue' maps to 'total_amount', 'customer name' maps to 'first_name + last_name'.

Stage 3 — Retriever: This is the RAG component. It embeds the question using OpenAI text-embedding-3-small and searches Qdrant for the top 4 most semantically similar past Q→SQL examples. These become the few-shot context for the LLM.

Stage 4 — Generator: Builds a structured prompt with the schema, the retrieved examples, and the user's question, then calls GPT-4o to generate SQL.

Stage 5 — Validator: Checks the SQL using the sqlglot AST parser for syntax errors, validates table and column names against the live schema, and enforces the guardrail rules — SELECT allowed, DELETE blocked.

Stage 6 — Executor: Runs the SQL safely against PostgreSQL. If it fails, the error message is automatically injected back into GPT-4o for one retry. I call this self-healing SQL.

Stage 7 — Formatter: Returns the results as a structured JSON table with column names, row count, and a natural language summary.

Each stage is an independent Python module — you can swap any one without touching the others."
""")


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — Key Features
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide(); bg(s)
slide_heading(s, "Key Features", "What makes QueryMind production-ready")

features = [
    (SKY,    "Few-Shot RAG",
     "Top-4 semantically similar Q→SQL examples retrieved from Qdrant at query time. "
     "30 curated examples covering 4 query types. Uses OpenAI text-embedding-3-small (1536-dim)."),
    (GREEN,  "Runtime Guardrails",
     "Toggle SELECT / INSERT / UPDATE / DELETE on or off at runtime. "
     "Max-rows limit (default 500). Enforced at the validator layer — never bypassed by clever prompts."),
    (RED,    "Self-Healing SQL",
     "On execution failure, the SQL and error message are automatically injected back into GPT-4o "
     "for one retry. Handles column mismatches and minor syntax issues without user intervention."),
    (YELLOW, "Full Audit Logging",
     "Every query logged: NL question, generated SQL, status (success/error/blocked), "
     "row count, and latency in ms. Filterable and paginated in the Admin UI."),
    (ORANGE, "Document Ingestion",
     "Upload PDF / DOCX / HTML schema documentation via Docling. "
     "Chunks are indexed into Qdrant schema_docs collection for context-aware SQL generation."),
    (PURPLE, "Runtime Model Config",
     "Switch between GPT-4o and GPT-4o-mini, change SQL dialect, adjust temperature — "
     "all through the Admin UI without restarting the server."),
]

for i, (color, title, desc) in enumerate(features):
    col = i % 2
    row = i // 2
    left = 0.4  + col * 6.5
    top  = 1.95 + row * 1.75
    box(s, left, top, 6.2, 1.6, fill_color=RGBColor(0x0D, 0x1F, 0x38))
    box(s, left, top, 0.1, 1.6, fill_color=color)
    txt(s, title, left + 0.25, top + 0.1, 5.8, 0.45,
        size=15, bold=True, color=color)
    txt(s, desc,  left + 0.25, top + 0.55, 5.7, 0.95,
        size=11, color=GRAY)

add_notes(s, """SLIDE 4 — KEY FEATURES (1 minute)

"Six features make QueryMind production-ready rather than just a demo.

Few-Shot RAG — instead of fine-tuning, I retrieve semantically similar past examples at query time. 30 curated examples across 4 query types are stored in Qdrant and retrieved by cosine similarity. This dramatically improves SQL accuracy.

Runtime Guardrails — each SQL operation type can be toggled on or off through the Admin UI. The max rows limit caps results at 500 by default. These are enforced at the validator stage — they cannot be bypassed by clever phrasing.

Self-Healing SQL — if the generated SQL fails on execution, the error is automatically fed back to GPT-4o for one retry. This handles most real-world failures like column name mismatches silently.

Full Audit Logging — every single query is recorded. You can see exactly what SQL was generated, whether it succeeded or was blocked, how many rows were returned, and how long it took.

Document Ingestion — you can upload schema documentation as PDF or DOCX. Docling parses it and indexes it into Qdrant so the LLM has richer context when generating SQL.

Runtime Config — switch the LLM model, dialect, or temperature through the Admin UI. No server restart needed."
""")


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — Admin Panel
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide(); bg(s)
slide_heading(s, "Admin Panel — 5 Management Screens",
              "Full React UI  •  Next.js 14 + Tailwind CSS  •  No code changes needed")

screens = [
    (SKY,    "Schema Manager",
     "/admin/schema",
     ["Browse all 7 live tables", "Add column descriptions & synonyms",
      "Mark sensitive columns", "Refresh from DB anytime"]),
    (YELLOW, "Examples Manager",
     "/admin/examples",
     ["Add / edit / delete Q→SQL pairs", "Filter by query type",
      "Sync all to Qdrant embeddings", "30 examples pre-seeded"]),
    (GREEN,  "Query Logs",
     "/admin/logs",
     ["Full audit trail of every query", "Filter: Success / Error / Blocked",
      "Expand to see generated SQL", "Pagination — 20 per page"]),
    (RED,    "Guardrails Config",
     "/admin/guardrails",
     ["Toggle SELECT / INSERT / UPDATE / DELETE", "Set max_rows limit",
      "Changes take effect immediately", "Blocked queries auto-logged"]),
    (PURPLE, "Model Config",
     "/admin/config",
     ["Switch GPT-4o / GPT-4o-mini", "Change SQL dialect",
      "Adjust temperature", "No server restart required"]),
]

for i, (color, name, url, bullets) in enumerate(screens):
    col = i % 3
    row = i // 3
    left = 0.35 + col * 4.35
    top  = 1.95 + row * 2.6
    box(s, left, top, 4.1, 2.4, fill_color=RGBColor(0x0D, 0x1F, 0x38))
    box(s, left, top, 4.1, 0.48, fill_color=color)
    txt(s, name, left + 0.12, top + 0.06, 3.86, 0.36,
        size=13, bold=True, color=DARK)
    txt(s, url,  left + 0.12, top + 0.56, 3.86, 0.3,
        size=10, color=color, italic=True)
    bullet_block(s, bullets,
                 left + 0.12, top + 0.88, 3.86, 1.4,
                 size=11, color=GRAY, bullet="• ")

add_notes(s, """SLIDE 5 — ADMIN PANEL (1 minute)

"The admin panel has five management screens, all built in React with Next.js 14 and Tailwind CSS.

Schema Manager — shows the live database schema pulled directly from PostgreSQL. You can add descriptions and synonyms to columns, which helps the LLM understand domain-specific terminology. For example, tagging 'total_amount' with the synonym 'revenue' means users can ask 'What is total revenue?' and the LLM correctly maps it.

Examples Manager — this is where you manage the few-shot Q→SQL pairs. You can add, edit, or delete examples, and trigger a full re-sync to Qdrant. I've pre-seeded 30 examples.

Query Logs — a full audit trail of every query. You can expand any row to see the exact SQL that was generated, filter by success, error, or blocked status, and see the latency.

Guardrails Config — the control panel for SQL permissions. Toggle operations on or off in real time. The changes are immediate — no restart needed.

Model Config — switch between GPT-4o and GPT-4o-mini, change the dialect, adjust temperature. All runtime, all through the UI."
""")


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — Evaluation Results
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide(); bg(s)
slide_heading(s, "Evaluation Results",
              "Automated harness  •  25 ground-truth queries  •  4 query types")

# Big metric cards
metrics = [
    ("100%", "Execution Accuracy", "25 / 25 queries\nran successfully", GREEN),
    ("8%",   "Exact Match Rate",   "Low due to LIMIT 500\nguardrail — expected", YELLOW),
    ("~3s",  "Avg Latency",        "End-to-end: classify\n→ generate → execute", SKY),
    ("30",   "Few-Shot Examples",  "Seeded across all\n4 query types", ORANGE),
]
for i, (val, label, note, color) in enumerate(metrics):
    left = 0.4 + i * 3.12
    box(s, left, 1.85, 2.85, 2.3, fill_color=RGBColor(0x0D, 0x1F, 0x38))
    txt(s, val,   left, 1.95, 2.85, 1.0,
        size=46, bold=True, color=color, align=PP_ALIGN.CENTER)
    txt(s, label, left, 2.9,  2.85, 0.45,
        size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s, note,  left, 3.38, 2.85, 0.65,
        size=10, color=GRAY, align=PP_ALIGN.CENTER)

# Per-type breakdown table
col_labels = ["Query Type", "Count", "Exec Accuracy", "Exact Match"]
col_x      = [0.4, 5.2, 7.0, 9.5]
col_w      = [4.6, 1.6, 2.3, 2.3]
rows_data  = [
    ["SELECT_SIMPLE",    "6", "100%", "16.7%"],
    ["SELECT_AGGREGATE", "7", "100%", "0.0%"],
    ["SELECT_JOIN",      "7", "100%", "14.3%"],
    ["SELECT_TEMPORAL",  "5", "100%", "0.0%"],
]

top_tbl = 4.4
# Header
for label_t, cx, cw in zip(col_labels, col_x, col_w):
    box(s, cx, top_tbl, cw - 0.05, 0.4, fill_color=SKY)
    txt(s, label_t, cx + 0.05, top_tbl + 0.06, cw - 0.1, 0.3,
        size=12, bold=True, color=DARK, align=PP_ALIGN.CENTER)

for ri, row in enumerate(rows_data):
    bg_c = RGBColor(0x0D, 0x1F, 0x38) if ri % 2 == 0 else RGBColor(0x0A, 0x18, 0x2E)
    for cell, cx, cw in zip(row, col_x, col_w):
        box(s, cx, top_tbl + 0.4 + ri * 0.4, cw - 0.05, 0.4, fill_color=bg_c)
        color = GREEN if cell == "100%" else WHITE
        txt(s, cell,
            cx + 0.05, top_tbl + 0.46 + ri * 0.4, cw - 0.1, 0.3,
            size=12, color=color, align=PP_ALIGN.CENTER)

txt(s, "Note: Exact match is low because the pipeline appends LIMIT 500 (max_rows guardrail). Results are correct — execution accuracy is the primary metric.",
    0.4, 7.0, 12.5, 0.38, size=10, color=GRAY, italic=True, align=PP_ALIGN.CENTER)

add_notes(s, """SLIDE 6 — EVALUATION RESULTS (45 seconds)

"I built a formal evaluation harness with 25 ground-truth question-to-SQL pairs covering all four query types — simple selects, aggregates, joins, and temporal queries.

The headline result is 100% execution accuracy — every single query ran successfully and returned correct data.

The exact match rate is 8% — and this is expected, not a failure. The pipeline automatically appends LIMIT 500 to every query as a guardrail. That differs from the bare ground-truth SQL, so it doesn't count as an exact match. But the actual SQL logic is correct.

Breaking it down by query type — all four categories achieved 100% execution accuracy. Simple, aggregate, join, and temporal — all working perfectly.

The evaluation harness is at eval/run_eval.py in the repo and can be re-run against any live backend instance at any time."
""")


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — Live Demo Queries
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide(); bg(s)
slide_heading(s, "Live Demo — Sample Queries",
              "Type these into the Query UI at localhost:3000/query")

demo_qs = [
    (SKY,    "SELECT_SIMPLE",    "List all pending orders"),
    (SKY,    "SELECT_SIMPLE",    "Show products that are out of stock"),
    (GREEN,  "SELECT_AGGREGATE", "How many customers do we have?"),
    (GREEN,  "SELECT_AGGREGATE", "What is the total revenue from delivered orders?"),
    (GREEN,  "SELECT_AGGREGATE", "How many orders are there per status?"),
    (PURPLE, "SELECT_JOIN",      "Show all orders with customer names"),
    (PURPLE, "SELECT_JOIN",      "List the top 5 customers by total amount spent"),
    (PURPLE, "SELECT_JOIN",      "Which products have been ordered more than 10 times?"),
    (YELLOW, "SELECT_TEMPORAL",  "Show orders placed in the last 30 days"),
    (YELLOW, "SELECT_TEMPORAL",  "Show monthly order counts for this year"),
]

for i, (color, qtype, question) in enumerate(demo_qs):
    col = i % 2
    row = i // 2
    left = 0.4  + col * 6.5
    top  = 1.85 + row * 1.08
    box(s, left, top, 6.1, 0.95, fill_color=RGBColor(0x0D, 0x1F, 0x38))
    box(s, left, top, 0.08, 0.95, fill_color=color)
    txt(s, qtype.replace("SELECT_", ""),
        left + 0.2, top + 0.06, 2.2, 0.3,
        size=9, color=color, bold=True)
    txt(s, f'"{question}"',
        left + 0.2, top + 0.48, 5.7, 0.4,
        size=13, color=WHITE, bold=True)

# Guardrail demo note
box(s, 0.4, 7.05, 12.5, 0.38,
    fill_color=RGBColor(0x1E, 0x29, 0x3B))
txt(s, "Guardrail test — also try:  \"Delete all customers\"  →  blocked  |  Turn off allow_select in Admin → all queries blocked",
    0.6, 7.09, 12.1, 0.3, size=10, color=RED, bold=True)

add_notes(s, """SLIDE 7 — LIVE DEMO QUERIES (1 minute reference card)

"Now let me show the live system. I have the Query UI open at localhost:3000.

I'll start with two simple queries — pending orders and out-of-stock products — to show basic filtering working correctly.

Then three aggregate queries — customer count, total revenue, and orders by status. Watch how the LLM generates GROUP BY and SUM correctly based on the few-shot examples.

Then three JOIN queries. The top 5 customers query is particularly impressive — it involves a JOIN, GROUP BY, ORDER BY, and LIMIT all generated correctly in one shot.

And finally two temporal queries using INTERVAL and EXTRACT — these are the hardest for LLMs to get right without good examples.

For each result, notice three things:
1. The generated SQL shown in the response
2. The result table with real data from the database
3. The natural language summary explaining the result

Then I'll demo the guardrail — I'll type 'Delete all customers' and show it getting blocked immediately."
""")


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — Summary & Links
# ═══════════════════════════════════════════════════════════════════════════════
s = add_slide(); bg(s)
box(s, 0, 0, 13.33, 0.5, fill_color=SKY)
box(s, 0, 7.0, 13.33, 0.5, fill_color=SKY)

txt(s, "QueryMind AI — Summary", 1, 0.6, 11, 0.75,
    size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

achievements = [
    "7-stage modular NL→SQL pipeline with self-healing SQL",
    "Qdrant vector RAG — 30 curated few-shot examples across 4 query types",
    "Runtime guardrails — configurable without code changes or server restart",
    "Full audit logging — every query tracked with SQL, latency, and status",
    "React admin panel — 5 management screens (Schema, Examples, Logs, Guardrails, Config)",
    "Evaluation harness — 100% execution accuracy on 25 ground-truth queries",
    "Docling document ingestion — schema docs indexed to Qdrant for richer LLM context",
]
bullet_block(s, achievements,
             1.5, 1.5, 10.3, 4.2,
             size=15, color=WHITE, bullet="✅  ", line_space=Pt(8))

# Divider
box(s, 0.5, 5.85, 12.33, 0.02, fill_color=RGBColor(0x1E, 0x3A, 0x5F))

# Links
links = [
    ("GitHub", "github.com/vijaymamilla/querymind"),
    ("Demo Video", "youtu.be/L4pgpwKwABo"),
    ("Backend API", "localhost:8000/docs"),
    ("Admin UI", "localhost:3000"),
]
for i, (label_t, url) in enumerate(links):
    left = 0.8 + i * 3.0
    txt(s, label_t, left, 6.02, 2.7, 0.35,
        size=12, bold=True, color=GRAY, align=PP_ALIGN.CENTER)
    txt(s, url, left, 6.4, 2.7, 0.4,
        size=12, color=SKY, align=PP_ALIGN.CENTER)

add_notes(s, """SLIDE 8 — SUMMARY & CLOSING (30 seconds)

"To summarize what QueryMind delivers:

A complete 7-stage Text-to-SQL pipeline with self-healing. Qdrant-powered RAG with 30 curated few-shot examples. Runtime guardrails that actually block dangerous operations at the validator layer. Full audit logging of every query. A React admin panel with five management screens. 100% execution accuracy on a formal 25-query evaluation. And Docling document ingestion for schema-aware SQL generation.

The full source code is on GitHub at github.com/vijaymamilla/querymind. The demo video is at youtu.be/L4pgpwKwABo.

Thank you for your time. I'm happy to answer any questions about the architecture, the pipeline design decisions, or the evaluation methodology."
""")


# ── Save ──────────────────────────────────────────────────────────────────────
out = "QueryMind_Presentation_v2.pptx"
prs.save(out)
print(f"\n✅  Saved: {out}")
print(f"   Slides: {len(prs.slides)}")
print("\n   Speaker notes are embedded — view in PowerPoint Presenter View")
print("   To export PDF: File → Export → PDF")
print("   To see notes:  View → Notes Page  or  use Presenter View")
